"""Step 6: the four-sheet spreadsheet, the Shortlist CSV and the terminal summary."""

from __future__ import annotations

import csv
import json
import statistics
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import TextIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from mmscan.config import Config
from mmscan.models import local_time
from mmscan.offers import load_current
from mmscan.scoring import DEAL, WORTH, Scored, percentile
from mmscan.store import Store

DISCLAIMER = "Cached data. Confirm price and seats on miles-and-more.com before acting."
STALE_HOURS = 48
GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
MILES_COLUMNS = {"Miles"}

SHORTLIST_COLUMNS = [
    "Verdict",
    "Change",
    "Date",
    "Origin",
    "Destination",
    "Cabin",
    "Miles",
    "Miles cost (EUR)",
    "Fees (EUR)",
    "Total cost (EUR)",
    "Cash fare (EUR)",
    "Saving vs cash (EUR)",
    "Value per mile (ct)",
    "Route percentile",
    "Destination percentile",
    "Seats",
    "Stops",
    "Carriers",
    "Flights",
    "Departs",
    "Arrives",
    "Flags",
    "Data age (h)",
    "Booking link",
]
ALL_COLUMNS = SHORTLIST_COLUMNS + ["Fee source", "Seats OK", "Direct", "Expired"]


def eligible(results: list[Scored], cfg: Config | None = None) -> list[Scored]:
    """Priced, live results; with `cfg` and shortlist.scope "focus", only those in the benchmark
    or watchlist focus."""
    out = [s for s in results if not s.fee_estimated and not s.expired and s.total_eur is not None]
    if cfg is not None and cfg.shortlist.scope == "focus":
        out = [
            s
            for s in out
            if cfg.in_focus(s.origin, s.origin_region, s.destination, s.destination_region, s.cabin)
        ]
    return out


def shortlist(results: list[Scored], cfg: Config | None = None) -> list[Scored]:
    """Eligible Deal and Worth a look rows."""
    return [s for s in eligible(results, cfg) if s.verdict in (DEAL, WORTH)]


def _age_hours(updated_at: str | None, now: datetime) -> float | None:
    if not updated_at:
        return None
    try:
        t = datetime.fromisoformat(updated_at.replace("Z", "+00:00"))
    except ValueError:
        return None
    if t.tzinfo is None:
        t = t.replace(tzinfo=UTC)
    return (now - t).total_seconds() / 3600


def _age_cell(updated_at: str | None, now: datetime) -> str | int | None:
    hours = _age_hours(updated_at, now)
    if hours is None:
        return None
    n = int(round(hours))
    return f"{n} (stale)" if hours > STALE_HOURS else n


def _r(x: float | None, nd: int = 2) -> float | None:
    return None if x is None else round(x, nd)


def row_values(s: Scored, cfg: Config, now: datetime) -> list:
    flags = [
        name for name, on in (("chart_rate", s.chart_rate), ("ceiling_good", s.ceiling_good)) if on
    ]
    miles_cost = None if s.miles is None else round(s.miles * cfg.mile_cost_eur_cents / 100, 2)
    return [
        s.verdict,
        s.change or None,
        s.date,
        s.origin,
        s.destination,
        s.cabin,
        s.miles,
        miles_cost,
        _r(s.fees_eur),
        _r(s.total_eur),
        _r(s.cash_fare_eur),
        _r(s.saving_eur),
        _r(s.value_cents),
        _r(s.route_pct, 1),
        _r(s.destination_pct, 1),
        s.seats if s.seats is not None else "unknown",
        s.stops,
        s.carriers,
        s.flight_numbers,
        local_time(s.departs_at or "") or None,
        local_time(s.arrives_at or "") or None,
        ", ".join(flags) or None,
        _age_cell(s.updated_at, now),
        s.booking_link,
    ]


def _write_table(ws, header: list[str], rows: list[list]) -> None:
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(len(rows) + 1, 1)}"
    for i, name in enumerate(header, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(10, min(len(name) + 2, 40))
        if name in MILES_COLUMNS:
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0"
    if header and header[0] == "Verdict":
        for cell in ws["A"][1:]:
            if cell.value == DEAL:
                cell.fill = GREEN
            elif cell.value == WORTH:
                cell.fill = AMBER


def route_stats(results: list[Scored]) -> list[list]:
    priced = [s for s in results if s.total_eur is not None and not s.expired]
    groups: dict[tuple, list[Scored]] = {}
    for s in priced:
        groups.setdefault(("route", f"{s.origin}-{s.destination}", s.cabin), []).append(s)
        groups.setdefault(
            ("destination", f"{s.origin_region}-{s.destination}", s.cabin), []
        ).append(s)
    out = []
    for (level, name, cabin), items in sorted(groups.items()):
        totals = [s.total_eur for s in items]
        fees = [s.fees_eur for s in items if not s.fee_estimated and s.fees_eur is not None]
        out.append(
            [
                level,
                name,
                cabin,
                len(totals),
                _r(min(totals)),
                _r(percentile(totals, 20)),
                _r(statistics.median(totals)),
                _r(max(totals)),
                _r(statistics.median(fees)) if fees else None,
            ]
        )
    return out


def run_log(
    store: Store, run_id: int, results: list[Scored], now: datetime, cfg: Config
) -> list[list]:
    run = store.run_info(run_id)
    stats = json.loads(run["stats_json"] or "{}") if run else {}
    warnings = json.loads(run["warnings_json"] or "[]") if run else []
    calls = store.conn.execute(
        "SELECT count(*) FROM api_calls WHERE run_id=?", (run_id,)
    ).fetchone()[0]
    sl = shortlist(results, cfg)
    rows = [
        ["Disclaimer", DISCLAIMER],
        ["Run id", run_id],
        ["Run started (UTC)", run["started_at"] if run else None],
        ["Report written (UTC)", now.isoformat(timespec="seconds")],
        ["Status", run["status"] if run else None],
        ["API calls", calls],
        ["Results", len(results)],
        [
            "Priced results",
            sum(1 for s in results if not s.fee_estimated and s.total_eur is not None),
        ],
        [
            "Estimated-fee results",
            sum(1 for s in results if s.fee_estimated and s.total_eur is not None),
        ],
        ["Expired results", sum(1 for s in results if s.expired)],
        ["Deal", sum(1 for s in sl if s.verdict == DEAL)],
        ["Worth a look", sum(1 for s in sl if s.verdict == WORTH)],
        ["Candidates not yet priced", stats.get("left_over", 0)],
    ]
    for k in sorted(stats):
        if k != "left_over":
            rows.append(
                [
                    f"stat: {k}",
                    json.dumps(stats[k]) if isinstance(stats[k], (dict, list)) else stats[k],
                ]
            )
    for w in warnings:
        rows.append(["Warning", w])
    rows.append(["Config", run["config_json"] if run else None])
    return rows


def write_report(store: Store, cfg: Config, run_id: int, now: datetime) -> list[Path]:
    """The current offers as a spreadsheet and CSV. `run_id` only fills the Run log sheet."""
    results = load_current(store)
    sl = shortlist(results, cfg)
    out_dir = Path(cfg.output.dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"mmscan_{now.strftime('%Y-%m-%d_%H%M')}"
    written: list[Path] = []

    if "xlsx" in cfg.output.formats:
        wb = Workbook()
        ws = wb.active
        ws.title = "Shortlist"
        _write_table(ws, SHORTLIST_COLUMNS, [row_values(s, cfg, now) for s in sl])
        all_rows = [
            row_values(s, cfg, now)
            + [
                "estimated" if s.fee_estimated else "actual",
                s.seats_ok,
                s.direct,
                s.expired,
            ]
            for s in results
        ]
        _write_table(wb.create_sheet("All results"), ALL_COLUMNS, all_rows)
        _write_table(
            wb.create_sheet("Route stats"),
            [
                "Level",
                "Group",
                "Cabin",
                "Samples",
                "Min (EUR)",
                "P20 (EUR)",
                "Median (EUR)",
                "Max (EUR)",
                "Median fee (EUR)",
            ],
            route_stats(results),
        )
        log_ws = wb.create_sheet("Run log")
        _write_table(log_ws, ["Item", "Value"], run_log(store, run_id, results, now, cfg))
        path = out_dir / f"{stem}.xlsx"
        wb.save(path)
        written.append(path)

    if "csv" in cfg.output.formats:
        path = out_dir / f"{stem}_shortlist.csv"
        with path.open("w", newline="") as f:
            w = csv.writer(f)
            w.writerow(SHORTLIST_COLUMNS)
            for s in sl:
                w.writerow(["" if v is None else v for v in row_values(s, cfg, now)])
        written.append(path)
    return written


def print_summary(
    store: Store, run_id: int, cfg: Config, now: datetime, file: TextIO = sys.stdout
) -> None:
    results = load_current(store)
    sl = shortlist(results, cfg)
    calls = store.conn.execute(
        "SELECT count(*) FROM api_calls WHERE run_id=?", (run_id,)
    ).fetchone()[0]
    deals = sum(1 for s in sl if s.verdict == DEAL)
    print(
        f"Run {run_id}: {len(results)} results, {len(sl)} on the shortlist "
        f"({deals} Deal, {len(sl) - deals} Worth a look), {calls} API calls.",
        file=file,
    )
    header = ["Verdict", "Date", "Route", "Cabin", "Miles", "Fees", "Total", "Seats", "Stops"]
    rows = [
        [
            s.verdict,
            s.date or "",
            f"{s.origin}-{s.destination}",
            s.cabin,
            f"{s.miles:,}" if s.miles is not None else "",
            f"{s.fees_eur:.0f}" if s.fees_eur is not None else "",
            f"{s.total_eur:.0f}" if s.total_eur is not None else "",
            str(s.seats) if s.seats is not None else "unknown",
            "" if s.stops is None else str(s.stops),
        ]
        for s in sl[:10]
    ]
    if rows:
        widths = [max(len(str(x)) for x in col) for col in zip(header, *rows, strict=False)]
        for r in [header, *rows]:
            print("  ".join(str(x).ljust(w) for x, w in zip(r, widths, strict=True)), file=file)
    else:
        print("No shortlist rows.", file=file)
    print(DISCLAIMER, file=file)
