from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from mmscan.scoring import score_run
from mmscan.store import Store
from mmscan.web.app import create_app
from mmscan.web.auth import MAX_FAILURES, LoginLimiter, password_ok
from mmscan.web.settings import SettingsError, WebSettings, settings_from_env
from tests.helpers import avail, config, trip

ROOT = Path(__file__).resolve().parents[1]
PASSWORD = "correct horse"
NOW = datetime.now(UTC)


@pytest.fixture
def db(tmp_path, default_config_dict):
    """A SQLite store with three scored scan runs on MUC-SIN business."""
    cfg = config(default_config_dict, relative={"percentile_good": 20, "min_samples": 2})
    path = tmp_path / "web.db"
    store = Store(path)
    for days, miles in ((2, 70000), (1, 64000), (0, 58000)):
        at = NOW - timedelta(days=days)
        run = store.start_run("scan", "{}", now=at)
        store.save_availability(run, avail(f"a{days}", business=miles), "watchlist")
        store.save_trips(
            f"a{days}",
            "2026-10-10T08:00:00Z",
            [trip(f"t{days}", avail_id=f"a{days}", miles=miles, taxes=30000)],
            "https://example.test/book",
            {},
        )
        score_run(store, cfg, run, at)
        store.record_call(run, "/search", {}, 200, 5, 1, 880 + days, now=at)
        store.finish_run(run, "partial" if days == 1 else "ok", {}, [], now=at)
    store.close()
    return path


@pytest.fixture
def client(db):
    settings = WebSettings(PASSWORD, "s" * 32, db, ROOT / "config.yaml", secure_cookies=True)
    return TestClient(create_app(settings), base_url="https://testserver")


def login(c: TestClient, password: str = PASSWORD):
    return c.post("/login", data={"password": password}, follow_redirects=False)


def test_healthz_is_public(client):
    assert client.get("/healthz").text == "ok"


@pytest.mark.parametrize(
    "path",
    [
        "/",
        "/shortlist",
        "/runs",
        "/trips",
        "/quota",
        "/route/MUC/SIN/business",
        "/reports/latest.xlsx",
        "/reports/latest-shortlist.csv",
    ],
)
def test_anonymous_requests_are_redirected_to_login(client, path):
    r = client.get(path, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/login")


def test_wrong_password_fails_and_right_one_works(client):
    r = login(client, "nope")
    assert r.status_code == 401
    assert client.get("/runs", follow_redirects=False).status_code == 303
    r = login(client)
    assert r.status_code == 303 and r.headers["location"] == "/"
    cookie = r.headers["set-cookie"].lower()
    assert "httponly" in cookie and "secure" in cookie and "samesite=lax" in cookie
    assert client.get("/runs").status_code == 200


def test_login_does_not_redirect_off_site(client):
    r = client.post(
        "/login", data={"password": PASSWORD, "next": "//evil.test"}, follow_redirects=False
    )
    assert r.headers["location"] == "/"


def test_logout_ends_the_session(client):
    login(client)
    client.post("/logout")
    assert client.get("/", follow_redirects=False).status_code == 303


def test_login_is_rate_limited(client):
    for _ in range(MAX_FAILURES):
        assert login(client, "x").status_code == 401
    r = login(client)  # even the right password waits
    assert r.status_code == 200 and "Too many attempts" in r.text


def test_limiter_window_and_compare():
    t = [0.0]
    lim = LoginLimiter(clock=lambda: t[0])
    for _ in range(MAX_FAILURES):
        lim.fail("c")
    assert lim.blocked("c")
    t[0] = 10_000
    assert not lim.blocked("c")
    assert password_ok("a", "a") and not password_ok("a", "b")


def test_shortlist_page_shows_fixture_data(client):
    login(client)
    r = client.get("/shortlist")
    assert r.status_code == 200
    assert "Cached data. Confirm price and seats on miles-and-more.com" in r.text
    assert "MUC→SIN" in r.text and "58,000" in r.text
    assert "https://example.test/book" in r.text
    # filters narrow the table
    assert "58,000" not in client.get("/shortlist?cabin=economy").text
    assert "Nothing matches" in client.get("/shortlist?route=FRA-KUL").text


def test_route_runs_quota_pages(client):
    login(client)
    r = client.get("/route/MUC/SIN/business")
    assert r.status_code == 200 and "<svg" in r.text and "Table view" in r.text
    runs = client.get("/runs").text
    assert 'class="partial"' in runs  # the partial run is highlighted
    q = client.get("/quota").text
    assert "of 300 calls used today" in q and "880" in q


def test_report_downloads(client):
    login(client)
    r = client.get("/reports/latest.xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert "Shortlist" in wb.sheetnames
    r = client.get("/reports/latest-shortlist.csv")
    assert r.status_code == 200 and r.text.startswith("Verdict,")


def test_no_endpoint_triggers_a_scan(client):
    login(client)
    for path in ("/scan", "/run", "/watch"):
        assert client.post(path).status_code in (404, 405)


def test_empty_database_renders(tmp_path):
    settings = WebSettings(PASSWORD, "s" * 32, tmp_path / "e.db", ROOT / "config.yaml")
    c = TestClient(create_app(settings), base_url="https://testserver")
    login(c)
    assert "No scored run yet" in c.get("/").text
    assert c.get("/reports/latest.xlsx").status_code == 404


def test_refuses_to_start_without_secrets(monkeypatch):
    monkeypatch.delenv("DASHBOARD_PASSWORD", raising=False)
    monkeypatch.setenv("SESSION_SECRET", "x")
    with pytest.raises(SettingsError, match="DASHBOARD_PASSWORD"):
        settings_from_env()
    monkeypatch.setenv("DASHBOARD_PASSWORD", "p")
    monkeypatch.delenv("SESSION_SECRET")
    with pytest.raises(SettingsError, match="SESSION_SECRET"):
        settings_from_env()


def test_trips_page_without_trips_explains_setup(client):
    login(client)
    r = client.get("/trips")
    assert r.status_code == 200 and "No trips configured" in r.text


def test_trips_page_lists_pairs(tmp_path, default_config_dict, write_config):
    from tests.test_trips import LEGS, TRIP, add_run
    from tests.test_trips import NOW as T_NOW

    cfg_path = write_config({**default_config_dict, "trips": [TRIP]})
    path = tmp_path / "trips.db"
    store = Store(path)
    add_run(store, config(default_config_dict, trips=[TRIP]), LEGS, at=T_NOW)
    store.close()
    c = TestClient(
        create_app(WebSettings(PASSWORD, "s" * 32, path, cfg_path, secure_cookies=True)),
        base_url="https://testserver",
    )
    login(c)
    text = c.get("/trips").text
    assert "Family stay" in text and "2027-06-10" in text and "2027-08-25" in text
    assert "Cheapest for 2" in text and "pairs" in text


def test_shortlist_labels_dynamic_prices(tmp_path, default_config_dict):
    cfg = config(default_config_dict, relative={"percentile_good": 20, "min_samples": 2})
    path = tmp_path / "mix.db"
    store = Store(path)
    run = store.start_run("scan", "{}", now=NOW)
    for aid, airlines in (("fx", "SQ"), ("dyn", "LH")):
        store.save_availability(run, avail(aid, airlines=airlines, business=60000), "watchlist")
        store.save_trips(
            aid,
            "2026-10-10T08:00:00Z",
            [trip(f"t{aid}", avail_id=aid, miles=60000, taxes=30000, carriers=airlines)],
            None,
            {},
        )
    score_run(store, cfg, run, NOW)
    store.close()
    settings = WebSettings(PASSWORD, "s" * 32, path, ROOT / "config.yaml")
    c = TestClient(create_app(settings), base_url="https://testserver")
    login(c)
    page = c.get("/shortlist?all=1").text
    assert "dynamic, verify" in page and ">chart<" in page
    assert ">Unranked</span>" in page  # the dynamic fare has no history yet
    assert ">Unranked</span>" not in c.get("/shortlist").text  # and never reaches the Shortlist


def test_deals_view_reads_the_store(db, default_config_dict):
    from mmscan.web import queries

    cfg = config(default_config_dict, relative={"percentile_good": 20, "min_samples": 2})
    store = Store(db)
    view = queries.deals_view(store, cfg, store.latest_scored_run(), NOW, pax=2, hide_stale=False)
    store.close()
    # one flight date seen three times is one observation: too little history for a baseline
    assert view["deals"] == [] and len(view["building"]) == 1
    assert view["thin_routes"] == 1 and view["total"] >= 1


@pytest.fixture
def deals_client(tmp_path, default_config_dict, write_config):
    """Eight flight dates on MUC-SIN business; the newest run holds the cheapest one."""
    cfg = config(default_config_dict, relative={"percentile_good": 20, "min_samples": 3})
    path = tmp_path / "deals.db"
    store = Store(path)
    at = NOW - timedelta(hours=1)
    run = store.start_run("scan", "{}", now=at)
    prices = [70000, 68000, 66000, 64000, 62000, 60000, 58000, 36000]
    for i, miles in enumerate(prices):
        aid = f"d{i}"
        store.save_availability(
            run,
            avail(aid, date=f"2026-12-1{i}", business=miles, updated=at.isoformat()),
            "watchlist",
        )
        store.save_trips(
            aid,
            at.isoformat(),
            [trip(f"t{i}", avail_id=aid, miles=miles, taxes=30000)],
            "https://example.test/book",
            {},
        )
    score_run(store, cfg, run, at)
    store.finish_run(run, "ok", {}, [], now=at)
    store.finish_run(store.start_run("watch", "{}", now=at), "ok", {}, [], now=at)
    store.close()
    cfg_path = write_config(
        {**default_config_dict, "relative": {"percentile_good": 20, "min_samples": 3}}
    )
    settings = WebSettings(PASSWORD, "s" * 32, path, cfg_path, secure_cookies=True)
    c = TestClient(create_app(settings), base_url="https://testserver")
    login(c)
    return c


def test_deals_page_ranks_against_the_usual_price(deals_client):
    r = deals_client.get("/?cabin=all&min_pct=0")
    assert r.status_code == 200
    assert "Best deals right now" in r.text
    assert "Great deal" in r.text and "Chart" in r.text
    assert "https://example.test/book" in r.text
    assert r.text.index("2026-12-17") < r.text.index("2026-12-16")  # biggest saving first


def test_deals_page_filters_and_traveller_scaling(deals_client):
    one = deals_client.get("/?cabin=all&pax=1&fresh=0&min_pct=0").text
    two = deals_client.get("/?cabin=all&pax=2&fresh=0&min_pct=0").text
    assert "36,000" in one and "72,000" in two
    assert "Nothing matches" in deals_client.get("/?cabin=economy").text
    assert "Nothing matches" in deals_client.get("/").text  # default cabin is economy
    assert "36,000" in deals_client.get("/?cabin=business&fresh=0").text
    assert "Nothing matches" in deals_client.get("/?cabin=all&route=FRA").text
    assert "Nothing matches" in deals_client.get("/?cabin=all&pax=9").text  # only 4 seats each


def test_deals_page_survives_bad_params(deals_client):
    assert deals_client.get("/?sort=bogus&pax=0").status_code == 200
    assert deals_client.get("/?pax=99").status_code == 200


def test_deals_page_empty_database(tmp_path):
    settings = WebSettings(PASSWORD, "s" * 32, tmp_path / "e.db", ROOT / "config.yaml")
    c = TestClient(create_app(settings), base_url="https://testserver")
    login(c)
    assert "No scored run yet" in c.get("/").text


def test_route_page_shows_usual_price_and_saving(deals_client):
    r = deals_client.get("/route/MUC/SIN/business")
    assert r.status_code == 200
    assert "Usual price" in r.text and "Best now vs usual" in r.text
    assert "Great deal" in r.text
    assert 'class="usual"' in r.text  # dashed line on the chart


def test_route_page_with_thin_history_says_so(db):
    settings = WebSettings(PASSWORD, "s" * 32, db, ROOT / "config.yaml", secure_cookies=True)
    c = TestClient(create_app(settings), base_url="https://testserver")
    login(c)
    r = c.get("/route/MUC/SIN/business")
    assert r.status_code == 200 and "building history" in r.text
    assert 'class="usual"' not in r.text


def test_price_chart_draws_usual_line():
    from mmscan.web.charts import price_chart

    pts = [{"ts": "2026-10-01T00:00:00Z", "low": 500.0, "median": 700.0, "n": 3}]
    assert 'class="usual"' in price_chart(pts, 900.0)
    assert 'class="usual"' not in price_chart(pts)


def _trips_client(tmp_path, default_config_dict, write_config, trip_extra=None, days=(0,)):
    from datetime import timedelta as td

    from tests.test_trips import LEGS, TRIP, add_run
    from tests.test_trips import NOW as T_NOW

    trip_cfg = {**TRIP, **(trip_extra or {})}
    cfg_path = write_config({**default_config_dict, "trips": [trip_cfg]})
    path = tmp_path / "trips2.db"
    store = Store(path)
    cfg = config(default_config_dict, trips=[trip_cfg])
    for back in days:  # days ago, oldest first
        add_run(store, cfg, LEGS, at=T_NOW - td(days=back))
    store.close()
    c = TestClient(
        create_app(WebSettings(PASSWORD, "s" * 32, path, cfg_path, secure_cookies=True)),
        base_url="https://testserver",
    )
    login(c)
    return c


def test_trips_page_calendar_trend_and_budget(tmp_path, default_config_dict, write_config):
    c = _trips_client(
        tmp_path, default_config_dict, write_config, {"max_total_eur": 9000}, days=(3, 2, 1, 0)
    )
    text = c.get("/trips").text
    assert "Cheapest day to leave" in text and 'class="cal"' in text
    assert "2027-06-10" in text and "no seats found" in text  # a priced day and an empty one
    assert '<svg class="spark"' in text and "Best price, last 30 days" in text
    assert "under the €9,000 budget" in text


def test_trips_page_without_history_says_so(tmp_path, default_config_dict, write_config):
    c = _trips_client(tmp_path, default_config_dict, write_config)
    text = c.get("/trips").text
    assert "Not enough history yet" in text and "Unverified" not in text


def test_trips_page_marks_dynamic_carriers(tmp_path, default_config_dict, write_config):
    d = {**default_config_dict, "dynamic_airlines": ["SQ"]}
    c = _trips_client(tmp_path, d, write_config)
    assert "Unverified" in c.get("/trips").text


def test_data_page(deals_client):
    text = deals_client.get("/data").text
    assert "Data &amp; alerts" in text and "API calls today" in text
    assert "How far to trust cached miles" in text and "not checked yet" in text  # LH etc.
    assert "Slack alerts" in text and "not sent" in text  # dynamic prices off by default
    assert "Healthy" in text or "Late" in text


def test_data_page_shows_manual_price_checks(db, default_config_dict):
    store = Store(db)
    store.add_price_check("lh", 6000, 17000, "MUC-SIN", "business")
    store.add_price_check("LH", 10000, 12000)
    store.close()
    c = TestClient(
        create_app(WebSettings(PASSWORD, "s" * 32, db, ROOT / "config.yaml", secure_cookies=True)),
        base_url="https://testserver",
    )
    login(c)
    text = c.get("/data").text
    assert "2.02×" in text  # median of 2.83 and 1.2
    assert "1.20× to 2.83×" in text


def test_usd_rate_warning_when_old(default_config_dict):
    from datetime import date

    from mmscan.scoring import _currency_warnings

    old = config(default_config_dict, usd_to_eur=0.92, usd_to_eur_date=date(2020, 1, 1))
    assert any("last checked" in w for w in _currency_warnings({"USD": 3}, old))
    fresh = config(default_config_dict, usd_to_eur=0.92, usd_to_eur_date=date.today())
    assert not any("last checked" in w for w in _currency_warnings({"USD": 3}, fresh))
    assert old.usd_rate_age_days(date(2020, 2, 1)) == 31


def test_all_routes_lists_every_flight_not_one_per_route(deals_client):
    r = deals_client.get("/overview?cabin=all&fresh=0&pax=1")
    assert r.status_code == 200
    assert r.text.count("MUC→SIN") == 8  # eight dates, one row each
    assert "Nothing matches" in deals_client.get("/overview?fresh=0").text  # default cabin economy
    assert (
        "Nothing matches"
        in deals_client.get("/overview?cabin=all&direction=Asia+%E2%86%92+Europe").text
    )


def test_best_deals_shows_only_big_savings_and_defaults_to_one_traveller(deals_client):
    big = deals_client.get("/?cabin=all&fresh=0").text
    assert "36,000" in big  # one traveller by default, not min_seats
    assert "60,000" not in big  # saves a few percent: under the 30% cut
    assert "60,000" in deals_client.get("/?cabin=all&fresh=0&min_pct=0").text


def test_hide_chart_deals_toggle(deals_client):
    assert "Nothing matches" in deals_client.get("/overview?cabin=all&fresh=0&nochart=1").text
    assert "MUC→SIN" in deals_client.get("/overview?cabin=all&fresh=0&chart=1").text
    both = deals_client.get("/overview?cabin=all&fresh=0&chart=1&nochart=1").text
    assert "Nothing matches" in both  # hide wins over chart-only


def test_every_row_links_to_miles_and_more(default_config_dict):
    from mmscan.config import load_config
    from mmscan.links import mm_link
    from tests.helpers import scored

    cfg = config(
        default_config_dict,
        miles_and_more=load_config(ROOT / "config.yaml").miles_and_more.model_dump(),
    )
    s = scored(400, origin="MUC", dest="FRA", date="2026-10-16", cabin="business")
    assert mm_link(cfg, s) == (
        "https://www.miles-and-more.com/row/en/spend/flights/flight-award.html?adults=1"
        "&bookingclass=B&children=0&dateformat=MM%2Fdd%2Fyyyy&destinationcode=FRA&infants=0"
        "&l=GB&origincode=MUC&outbounddate=10%2F16%2F2026&travellertype=M&triptype=OW&youths=0"
    )
    assert "adults=3" in mm_link(cfg, s, pax=3)
    assert "bookingclass=E" in mm_link(cfg, scored(400, cabin="economy"))
    assert mm_link(cfg, scored(400, cabin="premium"), 1) is None  # no known booking class
    s.booking_link = "https://stored.test/x"
    assert mm_link(cfg, s) == "https://stored.test/x"  # seats.aero's own link wins
    off = config(default_config_dict, miles_and_more={"link_template": ""})
    assert mm_link(off, scored(400)) is None  # no template


def test_link_template_rejects_unknown_placeholders(default_config_dict):
    import pytest
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        config(default_config_dict, miles_and_more={"link_template": "https://x/{nope}"})


def test_row_limit_and_show_more(deals_client):
    page = deals_client.get("/overview?cabin=all&fresh=0&pax=1&limit=50").text
    assert "Showing the first" not in page  # only 8 flights, under the limit
    r = deals_client.get("/overview?cabin=all&fresh=0&pax=1&limit=-5")
    assert r.status_code == 200  # a bad limit is clamped, not an error


def test_deals_pages_ignore_newer_watch_runs(db, default_config_dict):
    store = Store(db)
    run = store.start_run("watch", "{}", now=NOW + timedelta(hours=1))
    store.save_availability(run, avail("w1", business=50000), "watchlist")
    cfg = config(default_config_dict, relative={"percentile_good": 20, "min_samples": 2})
    score_run(store, cfg, run, NOW + timedelta(hours=1))
    store.finish_run(run, "ok", {}, [], now=NOW + timedelta(hours=1))
    store.close()
    c = TestClient(
        create_app(WebSettings(PASSWORD, "s" * 32, db, ROOT / "config.yaml", secure_cookies=True)),
        base_url="https://testserver",
    )
    login(c)
    text = c.get("/overview?cabin=all&fresh=0&pax=1").text
    assert f"Run #{run}" not in text
    assert "MUC→SIN" in text
