import csv
import io
from datetime import UTC, datetime

from openpyxl import load_workbook

from mmscan.report import DISCLAIMER, SHORTLIST_COLUMNS, print_summary, write_report
from mmscan.scoring import score_run
from mmscan.store import Store
from tests.helpers import avail, config, trip

NOW = datetime(2026, 10, 13, 12, 0, tzinfo=UTC)


def build(tmp_path, default_config_dict, n=12):
    cfg = config(
        default_config_dict,
        cabins=["business"],
        output={"dir": str(tmp_path / "reports"), "formats": ["xlsx", "csv"]},
    )
    store = Store(tmp_path / "db.sqlite")
    run = store.start_run("scan", "{}", now=NOW)
    for i in range(n):
        updated = "2026-10-13T08:00:00Z" if i else "2026-10-10T08:00:00Z"  # first one is stale
        store.save_availability(
            run,
            avail(f"a{i:02d}", date=f"2026-12-{i + 1:02d}", updated=updated, business=50000),
            "region",
        )
        store.save_trips(
            f"a{i:02d}",
            updated,
            [trip(f"t{i}", taxes=1000 if i % 3 == 0 else 40000 + i)],
            "https://book.test",
            {},
        )
    # One unpriced result on the same route gets an estimated fee.
    est = avail("est", date="2026-12-30", business=40000)
    del est["JTotalTaxes"]  # a fee of 0 is free; a missing one is unknown
    store.save_availability(run, est, "region")
    score_run(store, cfg, run, NOW)
    return store, cfg, run


def test_workbook_layout(tmp_path, default_config_dict):
    store, cfg, run = build(tmp_path, default_config_dict)
    xlsx, csv_path = write_report(store, cfg, run, NOW)
    assert xlsx.name == "mmscan_2026-10-13_1200.xlsx"
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Shortlist", "All results", "Route stats", "Run log"]
    ws = wb["Shortlist"]
    assert [c.value for c in ws[1]] == SHORTLIST_COLUMNS
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.startswith("A1:")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows, "expected shortlist rows"
    assert ws["A2"].value == "Deal" and ws["A2"].fill.fgColor.rgb.endswith("C6EFCE")
    assert ws["G2"].number_format == "#,##0"
    ages = [r[SHORTLIST_COLUMNS.index("Data age (h)")] for r in rows]
    assert "76 (stale)" in ages
    # The estimated-fee result is in All results but not on the Shortlist.
    all_rows = list(wb["All results"].iter_rows(min_row=2, values_only=True))
    assert any(r[-4] == "estimated" for r in all_rows)
    assert all(r[2] != "2026-12-30" for r in rows)
    log = {r[0]: r[1] for r in wb["Run log"].iter_rows(min_row=2, values_only=True)}
    assert log["Disclaimer"] == DISCLAIMER
    with csv_path.open() as f:
        assert next(csv.reader(f)) == SHORTLIST_COLUMNS


def test_empty_run_writes_all_sheets(tmp_path, default_config_dict):
    store, cfg, run = build(tmp_path, default_config_dict, n=0)
    xlsx, _ = write_report(store, cfg, run, NOW)
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Shortlist", "All results", "Route stats", "Run log"]
    assert wb["Shortlist"].max_row == 1


def test_summary_top_ten_and_disclaimer(tmp_path, default_config_dict):
    store, cfg, run = build(tmp_path, default_config_dict, n=15)
    out = io.StringIO()
    print_summary(store, run, cfg, NOW, out)
    lines = out.getvalue().splitlines()
    assert lines[-1] == DISCLAIMER
    table = [ln for ln in lines if ln.startswith(("Deal", "Worth a look"))]
    assert 0 < len(table) <= 10
